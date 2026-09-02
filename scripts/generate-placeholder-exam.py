#!/usr/bin/env python3
"""
Generates the placeholder TOEFL ITP final exam workbook (50 Listening / 40 Structure / 50 Reading).

This is SCAFFOLDING, not content. It exists so the exam -> scoring -> certificate -> verify chain
can be exercised end to end before the real syllabus arrives; every prompt is prefixed [CONTOH] so
a placeholder item can never be mistaken for something sold to a learner.

The real items replace these by re-uploading a sheet that keeps the same ids — that is exactly what
the importer's upsert is for. Run:  python3 scripts/generate-placeholder-exam.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

MARK = "[CONTOH]"
OUT = "backend/src/Infrastructure/Assets/PlaceholderFinalExam.xlsx"

COLUMNS = ["id", "section", "prompt", "choice_a", "choice_b", "choice_c", "choice_d",
           "answer", "passage_id", "audio_file", "tags"]

# ---------------------------------------------------------------- Listening (50)
# ITP part A: short conversations where the answer restates or infers, never repeats literally.
L_TEMPLATES = [
    ("M: I'm thinking of taking {topic} next term.\nW: Register early — it filled up in two days last year.\n"
     "Question: What does the woman imply?",
     ["The man should sign up quickly.", "Registration has already closed.",
      "The course is too hard for the man.", "She took the course last year."], "A"),
    ("W: Did you finish the {task}?\nM: Finish it? I haven't even started.\n"
     "Question: What does the man mean?",
     ["He is nowhere near finished.", "He submitted it yesterday.",
      "He lost his notes.", "He wants the woman's help."], "A"),
    ("M: Should we meet at the library or the {place}?\nW: Either works, but the library closes at six.\n"
     "Question: What does the woman suggest?",
     ["They should meet before six if it is the library.", "The library is already closed.",
      "She prefers not to meet today.", "The other place is too far."], "A"),
    ("W: How was the lecture on {topic}?\nM: Let's just say I've had more exciting afternoons.\n"
     "Question: What does the man mean?",
     ["He found the lecture dull.", "He enjoyed it very much.",
      "He missed the lecture.", "He arrived late."], "A"),
    ("M: I heard the {task} deadline moved.\nW: Moved? It's the same as it always was.\n"
     "Question: What does the woman say about the deadline?",
     ["It has not changed.", "It was brought forward.",
      "It was extended by a week.", "She does not know it."], "A"),
]
L_TOPICS = ["economics", "statistics", "linguistics", "geology", "psychology",
            "astronomy", "microbiology", "anthropology", "chemistry", "world history"]
L_TASKS = ["lab report", "research summary", "field notes", "problem set", "reading response",
           "case study", "annotated bibliography", "seminar paper", "data sheet", "book review"]
L_PLACES = ["student centre", "science building", "east cafeteria", "reading room", "media lab",
            "north annexe", "study hall", "computer suite", "seminar room", "quiet floor"]

# ---------------------------------------------------------------- Structure (40)
# ITP part B: 15 sentence-completion + 25 error-identification, mirroring the real paper's mix.
S_COMPLETION = [
    ("The committee ____ its decision at the end of the session.",
     ["announce", "announcing", "announced", "to announce"], "C"),
    ("Not until the results were published ____ the significance of the finding.",
     ["the team realised", "did the team realise", "the team did realise", "realised the team"], "B"),
    ("____ in the nineteenth century, the bridge still carries traffic today.",
     ["Built", "Building", "It was built", "To build"], "A"),
    ("The samples were divided into three groups, ____ was tested separately.",
     ["each of which", "each of them", "each", "which each"], "A"),
    ("Rarely ____ such a rapid change in a single decade.",
     ["scientists have observed", "have scientists observed",
      "scientists observed", "observed scientists"], "B"),
]
S_ERROR = [
    ("The number of applicants (A) have risen (B) sharply since (C) the programme began (D).\n"
     "Question: Which underlined part is incorrect?", ["A", "B", "C", "D"], "B"),
    ("Neither the students (A) nor the lecturer (B) were (C) aware of the change (D).\n"
     "Question: Which underlined part is incorrect?", ["A", "B", "C", "D"], "C"),
    ("The data collected (A) last year (B) suggests that (C) the trend are stable (D).\n"
     "Question: Which underlined part is incorrect?", ["A", "B", "C", "D"], "D"),
    ("Despite of (A) the delay, the team completed (B) the survey (C) on schedule (D).\n"
     "Question: Which underlined part is incorrect?", ["A", "B", "C", "D"], "A"),
    ("Each of the samples (A) were (B) examined under (C) a microscope (D).\n"
     "Question: Which underlined part is incorrect?", ["A", "B", "C", "D"], "B"),
]

# ---------------------------------------------------------------- Reading (50)
# Five passages, ten items each, covering main idea / detail / vocabulary / inference / purpose.
PASSAGES = [
    ("P1", "Early cities rarely grew in isolation. Where trade routes met, surplus grain and craft "
     "goods changed hands, and the settlements that controlled those crossings grew faster than "
     "their neighbours. Control of a crossing meant control of what moved through it, and the "
     "wealth that followed paid for walls, granaries and record-keeping."),
    ("P2", "The migration of monarch butterflies spans several generations. No single insect makes "
     "the entire round trip; instead, successive generations continue a journey none of them began. "
     "How the route is retained without individual memory remains an open question."),
    ("P3", "Glass is neither a conventional solid nor a liquid. Cooled quickly enough, its molecules "
     "are frozen before they can arrange into a crystal, leaving a disordered structure that behaves "
     "rigidly at ordinary temperatures. This accident of cooling explains both its transparency and "
     "its brittleness."),
    ("P4", "Before standardised time, each town kept its own noon by the sun. Railways made the "
     "practice untenable: a timetable spanning many towns needed one clock, not fifty. The zones "
     "adopted in the 1880s were a commercial convenience long before they were a legal fact."),
    ("P5", "Coral reefs occupy a fraction of the ocean floor yet shelter a quarter of marine species. "
     "The corals themselves depend on algae living within their tissues; when stress expels the "
     "algae, the coral pales and, if the condition persists, starves."),
]
R_TEMPLATES = [
    ("What is the main idea of the passage?",
     ["It is stated in the opening sentences.", "Trade was unimportant.",
      "The topic is purely modern.", "No conclusion is offered."], "A"),
    ("According to the passage, which statement is supported?",
     ["The one the passage states directly.", "A claim the passage contradicts.",
      "A detail never mentioned.", "An unrelated assertion."], "A"),
    ("The word \"surplus\" in the passage is closest in meaning to",
     ["extra", "spoiled", "imported", "hidden"], "A"),
    ("It can be inferred from the passage that",
     ["the outcome followed from the conditions described.", "the process was instantaneous.",
      "the author disputes the evidence.", "the topic has no modern relevance."], "A"),
    ("The author mentions the detail in order to",
     ["support the point being made.", "contradict the opening claim.",
      "introduce an unrelated topic.", "question the sources."], "A"),
    ("Which of the following is NOT mentioned in the passage?",
     ["A claim the passage never makes.", "A detail the passage states.",
      "A consequence the passage describes.", "A condition the passage names."], "A"),
    ("The passage suggests that the process described was",
     ["gradual rather than sudden.", "entirely accidental.",
      "reversed within a decade.", "confined to one location."], "A"),
    ("The word \"retained\" in the passage is closest in meaning to",
     ["kept", "lost", "measured", "shortened"], "A"),
    ("What does the passage imply about the subject's future?",
     ["It depends on the conditions the passage identifies.", "It is entirely settled.",
      "It is unrelated to the causes given.", "It was resolved long ago."], "A"),
    ("The passage is most likely taken from a text about",
     ["the academic subject it describes.", "personal correspondence.",
      "a work of fiction.", "a legal statute."], "A"),
]


def rotate(choices, answer, k):
    """
    Moves the correct choice to a different position, deterministically.

    Without this every Listening and Reading template answers "A", so answering A throughout
    would score about 79% — above any pass mark. A placeholder exam that cannot be failed is
    useless for testing the half of the product that handles failure.

    Structure items are deliberately NOT rotated: their choices are the literal labels A-D naming
    an underlined part of the sentence, so reordering them would list the options as B, C, D, A.
    Those templates already vary which part is wrong.
    """
    c = "ABCD".index(answer)
    rotated = [choices[(j - k) % 4] for j in range(4)]
    return rotated, "ABCD"[(c + k) % 4]


def rows():
    out = []
    for n in range(1, 51):
        prompt, choices, answer = L_TEMPLATES[(n - 1) % len(L_TEMPLATES)]
        prompt = prompt.format(topic=L_TOPICS[(n - 1) % len(L_TOPICS)],
                               task=L_TASKS[(n - 1) % len(L_TASKS)],
                               place=L_PLACES[(n - 1) % len(L_PLACES)])
        choices, answer = rotate(choices, answer, n % 4)
        out.append([f"PH-L{n:02}", "Listening", f"{MARK} {prompt}", *choices, answer,
                    "", "", "placeholder,itp,listening"])

    for n in range(1, 41):
        if n <= 15:
            prompt, choices, answer = S_COMPLETION[(n - 1) % len(S_COMPLETION)]
        else:
            prompt, choices, answer = S_ERROR[(n - 16) % len(S_ERROR)]
        out.append([f"PH-S{n:02}", "Structure", f"{MARK} {prompt}", *choices, answer,
                    "", "", "placeholder,itp,structure"])

    for n in range(1, 51):
        pid = PASSAGES[(n - 1) // 10][0]
        prompt, choices, answer = R_TEMPLATES[(n - 1) % len(R_TEMPLATES)]
        choices, answer = rotate(choices, answer, (n + 2) % 4)
        out.append([f"PH-R{n:02}", "Reading", f"{MARK} {prompt}", *choices, answer,
                    pid, "", "placeholder,itp,reading"])
    return out


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    for c, name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=c).value = name
        ws.cell(row=1, column=c).font = Font(bold=True)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["C"].width = 70
    for cell in ws["C"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].number_format = "@"   # keep ids as text

    data = rows()
    for r, row in enumerate(data, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c).value = value

    ps = wb.create_sheet("Passages")
    ps.cell(row=1, column=1).value = "passage_id"
    ps.cell(row=1, column=2).value = "text"
    ps["A1"].font = ps["B1"].font = Font(bold=True)
    ps.column_dimensions["B"].width = 100
    for r, (pid, text) in enumerate(PASSAGES, start=2):
        ps.cell(row=r, column=1).value = pid
        ps.cell(row=r, column=2).value = text
        ps.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 110
    lines = [
        "PLACEHOLDER CONTENT — NOT FOR SALE",
        "",
        f"Every prompt begins with {MARK} so a placeholder item can never be mistaken for real "
        "syllabus content, in the question bank or in front of a learner.",
        "",
        "Why it exists: the final exam, its timer, proctoring, ITP scoring, the certificate and the "
        "public verify page could not be tested end to end while the bank held 16 of the 140 items "
        "the format requires. This fills the shape so that chain is exercisable.",
        "",
        "How to replace it with real content: keep the ids (PH-L01, PH-S01, PH-R01 …), replace the "
        "prompts, choices and answers, and upload the sheet at /admin/questions/import. The importer "
        "upserts on id, so the real items overwrite these in place — no duplicates, no cleanup.",
        "",
        "Shape: 50 Listening, 40 Structure, 50 Reading = 140, the TOEFL ITP layout the readiness "
        "check enforces. Reading items reference five passages on the Passages sheet.",
        "",
        "Regenerate with: python3 scripts/generate-placeholder-exam.py",
    ]
    for i, line in enumerate(lines, start=1):
        info.cell(row=i, column=1).value = line
        info.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    info["A1"].font = Font(bold=True)

    wb.save(OUT)
    print(f"wrote {OUT}: {len(data)} questions, {len(PASSAGES)} passages")


if __name__ == "__main__":
    main()
