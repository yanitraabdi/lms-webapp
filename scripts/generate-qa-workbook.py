#!/usr/bin/env python3
"""
Builds docs/INVERTA_UAT_Test_Cases_v2.0.xlsx — the QA workbook, split into API and UI sheets.

Follows .claude/skills/writing-inverta-qa-test-cases/SKILL.md. Run:
    python3 scripts/generate-qa-workbook.py

v1.0 cases are carried over: those that are really API checks were rewritten with method, URL,
headers, body, one expected status and a runnable curl; the rest stayed UI and kept their wording.
Final-assessment cases are new — that chain only became testable once the 140-item placeholder
exam was seeded.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SRC = "docs/INVERTA_UAT_Test_Cases_v1.0.xlsx"
OUT = "docs/INVERTA_UAT_Test_Cases_v2.0.xlsx"
BASE = "https://inverta.recosta.id"

HEAD = PatternFill("solid", fgColor="1F2937")
HEADFONT = Font(bold=True, color="FFFFFF", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

# IDs from v1.0 that are really API checks — they move to the API sheet, rewritten.
MOVED_TO_API = {"A-03", "E-02", "K-01", "N-01", "N-02", "N-03", "N-05",
                "N-06", "N-07", "N-08", "N-15", "N-16", "N-19", "N-22", "H-02"}

# Rewritten below rather than carried over. E-05 was unrunnable on the only seeded learner
# (passing is permanent, GR-7 retains attempts) and produced two false bug reports; F-01 asked
# for an audio control that did not exist until the gating-audio endpoint was built.
REWRITTEN = {"E-05", "F-01"}

API_COLS = ["ID", "Description (EN)", "Deskripsi (ID)", "Preconditions (EN)", "Prasyarat (ID)",
            "Method + URL", "Headers", "Body", "Expected status", "Expected body",
            "Ready-to-run (curl)", "Result", "Tester", "Notes / Catatan"]

UI_COLS = ["ID", "Feature / Fitur", "Description (EN)", "Deskripsi (ID)",
           "Preconditions (EN)", "Prasyarat (ID)", "Steps (EN)", "Langkah (ID)",
           "Expected Result (EN)", "Hasil yang Diharapkan (ID)",
           "Result", "Tester", "Notes / Catatan"]

TOKEN = "Authorization: Bearer <token>"
JSON = "content-type: application/json"


def api(idv, en, idn, pre_en, pre_id, mu, headers, body, status, expect, curl):
    return dict(zip(API_COLS, [idv, en, idn, pre_en, pre_id, mu, headers, body,
                               status, expect, curl, "", "", ""]))


API_CASES = [
    # ---- auth ----
    api("API-AUTH-01",
        "Registering with an email that already exists is refused.",
        "Mendaftar dengan email yang sudah ada ditolak.",
        "siswa@test.local already exists (it is seeded).",
        "siswa@test.local sudah ada (sudah di-seed).",
        f"POST {BASE}/api/auth/register", JSON,
        '{"name":"QA Dup","email":"siswa@test.local","password":"Password123"}',
        "409",
        'Response contains "sudah terdaftar". No new account is created.',
        f"""curl -i -X POST {BASE}/api/auth/register -H '{JSON}' \\
  -d '{{"name":"QA Dup","email":"siswa@test.local","password":"Password123"}}'"""),

    api("API-AUTH-02",
        "A password shorter than 8 characters is refused.",
        "Kata sandi kurang dari 8 karakter ditolak.",
        "None. Use an email nobody has used.",
        "Tidak ada. Pakai email yang belum pernah dipakai.",
        f"POST {BASE}/api/auth/register", JSON,
        '{"name":"QA Short","email":"qa-short-01@test.local","password":"abc"}',
        "400",
        'errors.Password contains "must be at least 8 characters". No account is created.',
        f"""curl -i -X POST {BASE}/api/auth/register -H '{JSON}' \\
  -d '{{"name":"QA Short","email":"qa-short-01@test.local","password":"abc"}}'"""),

    api("API-AUTH-03",
        "Signing in with a wrong password fails without revealing whether the email exists.",
        "Masuk dengan kata sandi salah gagal tanpa mengungkap apakah email itu ada.",
        "None. Run it twice: once with siswa@test.local (exists), once with "
        "nobody-here@test.local (does not exist).",
        "Tidak ada. Jalankan dua kali: dengan siswa@test.local (ada), lalu "
        "nobody-here@test.local (tidak ada).",
        f"POST {BASE}/api/auth/login", JSON,
        '{"email":"siswa@test.local","password":"WrongPassword123"}',
        "401",
        "BOTH runs return the SAME status and the SAME message. If the existing email gives a "
        "different message, that leaks which emails are registered — report it as SECURITY.",
        f"""curl -i -X POST {BASE}/api/auth/login -H '{JSON}' \\
  -d '{{"email":"siswa@test.local","password":"WrongPassword123"}}'"""),

    api("API-AUTH-04",
        "An account is locked after 5 wrong passwords, and stays locked for the right one.",
        "Akun terkunci setelah 5 kata sandi salah, dan tetap terkunci untuk kata sandi yang benar.",
        "Use an account YOU registered, NOT siswa@test.local — this locks it for 15 minutes.",
        "Pakai akun yang ANDA daftarkan, BUKAN siswa@test.local — ini mengunci akun 15 menit.",
        f"POST {BASE}/api/auth/login", JSON,
        '{"email":"<your own email>","password":"WrongPassword123"}',
        "401",
        "Send the wrong password 5 times, then send the CORRECT password once. The 6th call is "
        "still refused. Lock lasts 15 minutes.",
        f"""for i in 1 2 3 4 5; do curl -s -o /dev/null -w "%{{http_code}}\\n" \\
  -X POST {BASE}/api/auth/login -H '{JSON}' \\
  -d '{{"email":"<your own email>","password":"WrongPassword123"}}'; done"""),

    # ---- access control (SECURITY) ----
    api("API-SEC-01",
        "SECURITY: a learner cannot call admin endpoints.",
        "KEAMANAN: pelajar tidak dapat memanggil endpoint admin.",
        "A learner token for siswa@test.local — see the Info sheet for how to get one.",
        "Token pelajar untuk siswa@test.local — lihat sheet Info untuk cara mendapatkannya.",
        f"GET {BASE}/api/admin/questions", TOKEN, "none",
        "403",
        "No question data is returned. Report immediately if this returns 200.",
        f"curl -i {BASE}/api/admin/questions -H '{TOKEN}'"),

    api("API-SEC-02",
        "SECURITY: admin endpoints cannot be reached without signing in at all.",
        "KEAMANAN: endpoint admin tidak dapat diakses tanpa masuk sama sekali.",
        "None. Send no Authorization header.",
        "Tidak ada. Jangan kirim header Authorization.",
        f"GET {BASE}/api/admin/questions", "none", "none",
        "401",
        "No question data is returned. 401 (not signed in), not 403.",
        f"curl -i {BASE}/api/admin/questions"),

    api("API-SEC-03",
        "SECURITY: the correct answers are never sent to the browser.",
        "KEAMANAN: kunci jawaban tidak pernah dikirim ke peramban.",
        "A learner token, and the id of a session that has a test. Get it from "
        "GET /api/me/programs/{programId}.",
        "Token pelajar, dan id sesi yang punya tes. Ambil dari "
        "GET /api/me/programs/{programId}.",
        f"GET {BASE}/api/sessions/{{sessionId}}/assessment", TOKEN, "none",
        "200",
        'The WHOLE response contains no "correct", no "answerKey", no "isCorrect". The grep '
        "below must print nothing. Any match is a SECURITY defect — report immediately.",
        f"""curl -s {BASE}/api/sessions/{{sessionId}}/assessment -H '{TOKEN}' \\
  | grep -io 'correct\\|answerkey\\|iscorrect'   # expect NO output"""),

    api("API-SEC-04",
        "SECURITY: a tampered audio link does not work.",
        "KEAMANAN: tautan audio yang diubah tidak berfungsi.",
        "Get a working signed audio URL from a Listening question first (API-EXAM-06). "
        "Then change one character in its sig= value.",
        "Dapatkan URL audio bertanda tangan yang berfungsi dari soal Listening (API-EXAM-06). "
        "Lalu ubah satu karakter pada nilai sig=.",
        f"GET {BASE}/api/media/audio/<key>?exp=<exp>&sig=<TAMPERED>", "none", "none",
        "404",
        "No audio is returned. 404 on purpose — a 403 would confirm the file exists.",
        f"curl -i '{BASE}/api/media/audio/<key>?exp=<exp>&sig=<TAMPERED>'"),

    # ---- enrollment ----
    api("API-ENR-01",
        "Starting checkout alone does NOT grant access. Critical rule.",
        "Memulai checkout saja TIDAK memberi akses. Aturan kritis.",
        "A token for a brand-new account that has never paid. Get the programId from "
        "GET /api/programs/toefl-preparation.",
        "Token untuk akun baru yang belum pernah membayar. Ambil programId dari "
        "GET /api/programs/toefl-preparation.",
        f"POST {BASE}/api/programs/{{programId}}/enroll  then  "
        f"GET {BASE}/api/me/programs/{{programId}}",
        f"{TOKEN}\n{JSON}", "{}",
        "403",
        "The enroll call returns a checkout URL, but the SECOND call (me/programs) must still "
        "be 403. Access is granted only by the paid webhook.",
        f"""curl -s -X POST {BASE}/api/programs/{{programId}}/enroll -H '{TOKEN}' -H '{JSON}' -d '{{}}'
curl -i {BASE}/api/me/programs/{{programId}} -H '{TOKEN}'   # must be 403"""),

    api("API-ENR-02",
        "A FAILED payment does not grant access.",
        "Pembayaran GAGAL tidak memberi akses.",
        "Run API-ENR-01 first and keep the providerRef from the enroll response.",
        "Jalankan API-ENR-01 dulu dan simpan providerRef dari respons enroll.",
        f"POST {BASE}/api/dev/payments/{{providerRef}}/fail  then  "
        f"GET {BASE}/api/me/programs/{{programId}}",
        TOKEN, "none",
        "403",
        "The second call is still 403. The learner has no access after a failed payment.",
        f"""curl -s -X POST {BASE}/api/dev/payments/{{providerRef}}/fail
curl -i {BASE}/api/me/programs/{{programId}} -H '{TOKEN}'   # must be 403"""),

    api("API-ENR-03",
        "An already-enrolled learner cannot be granted the same programme twice.",
        "Pelajar yang sudah terdaftar tidak dapat diberi program yang sama dua kali.",
        "An admin token, and siswa@test.local with ACTIVE access (restore it first via "
        "UI-ADM-08 if it is revoked).",
        "Token admin, dan siswa@test.local dengan akses AKTIF (pulihkan lewat UI-ADM-08 "
        "jika sedang dicabut).",
        f"POST {BASE}/api/admin/enrollments/grant",
        f"{TOKEN} (admin)\n{JSON}",
        '{"email":"siswa@test.local","programId":"<programId>"}',
        "409",
        'Response contains "sudah terdaftar aktif". No second enrollment row is created.',
        f"""curl -i -X POST {BASE}/api/admin/enrollments/grant -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"email":"siswa@test.local","programId":"<programId>"}}'"""),

    # ---- question bank validation ----
    api("API-QB-01",
        "A question with no correct answer marked is refused.",
        "Soal tanpa jawaban benar ditolak.",
        "An admin token. This state cannot be produced in the admin form — the radio group "
        "always keeps one choice selected — so it is only reachable here.",
        "Token admin. Kondisi ini tidak bisa dibuat lewat form admin — grup radio selalu "
        "menyisakan satu pilihan tertandai — jadi hanya bisa diuji di sini.",
        f"POST {BASE}/api/admin/questions", f"{TOKEN} (admin)\n{JSON}",
        '{"section":"Reading","prompt":"Q","choices":["a","b"],"correct":[],'
        '"audioRef":null,"passageRef":null,"tags":null}',
        "400",
        'Response contains "Tandai minimal satu jawaban benar." The question is not created.',
        f"""curl -i -X POST {BASE}/api/admin/questions -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"section":"Reading","prompt":"Q","choices":["a","b"],"correct":[],"audioRef":null,"passageRef":null,"tags":null}}'"""),

    api("API-QB-02",
        "A question with fewer than two answer choices is refused.",
        "Soal dengan kurang dari dua pilihan jawaban ditolak.",
        "An admin token. The admin form hides the delete button at two choices, so this is "
        "only reachable here.",
        "Token admin. Form admin menyembunyikan tombol hapus pada dua pilihan, jadi ini "
        "hanya bisa diuji di sini.",
        f"POST {BASE}/api/admin/questions", f"{TOKEN} (admin)\n{JSON}",
        '{"section":"Reading","prompt":"Q","choices":["hanya satu"],"correct":[0],'
        '"audioRef":null,"passageRef":null,"tags":null}',
        "400",
        'Response contains "Minimal 2 pilihan jawaban." The question is not created.',
        f"""curl -i -X POST {BASE}/api/admin/questions -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"section":"Reading","prompt":"Q","choices":["hanya satu"],"correct":[0],"audioRef":null,"passageRef":null,"tags":null}}'"""),

    api("API-QB-03",
        "A retake limit of 0 is refused, because it would lock the test forever.",
        "Batas percobaan 0 ditolak, karena akan mengunci tes selamanya.",
        "An admin token and the id of any assessment (GET /api/admin/assessments).",
        "Token admin dan id salah satu tes (GET /api/admin/assessments).",
        f"PUT {BASE}/api/admin/assessments/{{assessmentId}}", f"{TOKEN} (admin)\n{JSON}",
        '{"kind":"Gating","title":"Tes sesi","config":{"retakeCap":0}}',
        "400",
        'Response contains "Batas percobaan minimal 1." The test keeps its previous setting.',
        f"""curl -i -X PUT {BASE}/api/admin/assessments/{{assessmentId}} -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"kind":"Gating","title":"Tes sesi","config":{{"retakeCap":0}}}}'"""),

    api("API-QB-04",
        "A test that learners have already taken cannot be deleted.",
        "Tes yang sudah dikerjakan pelajar tidak dapat dihapus.",
        "An admin token, and the id of the session-1 gating test that siswa@test.local "
        "has already attempted.",
        "Token admin, dan id tes sesi 1 yang sudah dikerjakan siswa@test.local.",
        f"DELETE {BASE}/api/admin/assessments/{{assessmentId}}", f"{TOKEN} (admin)", "none",
        "409",
        "The test still exists afterwards. Learner history is never deleted.",
        f"curl -i -X DELETE {BASE}/api/admin/assessments/{{assessmentId}} -H '{TOKEN}'"),

    api("API-QB-05",
        "A score conversion table with a gap is reported as not ready.",
        "Tabel konversi skor yang bolong dilaporkan belum siap.",
        "An admin token and the programId. Do NOT run this on the live programme — use "
        "'QA Test Program'.",
        "Token admin dan programId. JANGAN jalankan pada program live — pakai "
        "'QA Test Program'.",
        f"GET {BASE}/api/admin/programs/{{programId}}/readiness", f"{TOKEN} (admin)", "none",
        "200",
        'A check with key "score_bands_complete" has "passed": false when a raw score in the '
        "range has no row. With the live seeded table it passes.",
        f"curl -s {BASE}/api/admin/programs/{{programId}}/readiness -H '{TOKEN}'"),

    # ---- final assessment ----
    api("API-EXAM-01",
        "Starting the final assessment without signing in is refused.",
        "Memulai tes akhir tanpa masuk ditolak.",
        "The final assessment id. Get it from GET /api/sessions/{finalSessionId}/assessment "
        "while signed in, then sign out for this call.",
        "Id tes akhir. Ambil dari GET /api/sessions/{finalSessionId}/assessment saat masuk, "
        "lalu keluar untuk panggilan ini.",
        f"POST {BASE}/api/assessments/{{assessmentId}}/attempts", "none", "none",
        "401",
        "No attempt is created.",
        f"curl -i -X POST {BASE}/api/assessments/{{assessmentId}}/attempts"),

    api("API-EXAM-02",
        "The sitting state reports the current section, its deadline and the strike count.",
        "Status pengerjaan melaporkan bagian saat ini, tenggat waktunya, dan jumlah peringatan.",
        "Run UI-EXAM-02 first to start the exam, and note the attemptId it creates.",
        "Jalankan UI-EXAM-02 dulu untuk memulai tes, dan catat attemptId yang dibuat.",
        f"GET {BASE}/api/attempts/{{attemptId}}/state", TOKEN, "none",
        "200",
        '"currentSection" is "Listening", a "deadline" timestamp is present, "strikes" is 0.',
        f"curl -s {BASE}/api/attempts/{{attemptId}}/state -H '{TOKEN}'"),

    api("API-EXAM-03",
        "The deadline is set by the server, so it does not reset when the page reloads.",
        "Tenggat waktu ditentukan server, jadi tidak ter-reset saat halaman dimuat ulang.",
        "Run API-EXAM-02 first and write down the exact \"deadline\" value.",
        "Jalankan API-EXAM-02 dulu dan catat nilai \"deadline\" persisnya.",
        f"GET {BASE}/api/attempts/{{attemptId}}/state", TOKEN, "none",
        "200",
        'Wait 60 seconds, call again: "deadline" is the SAME timestamp as before. It does not '
        "move forward. A changing deadline would mean the client can extend its own exam.",
        f"curl -s {BASE}/api/attempts/{{attemptId}}/state -H '{TOKEN}'   # run twice, 60s apart"),

    api("API-EXAM-04",
        "A blur shorter than 2 seconds is recorded but is never a strike.",
        "Blur di bawah 2 detik tercatat tetapi tidak pernah jadi peringatan.",
        "Run UI-EXAM-02 first. Use its attemptId.",
        "Jalankan UI-EXAM-02 dulu. Pakai attemptId-nya.",
        f"POST {BASE}/api/attempts/{{attemptId}}/proctor-events", f"{TOKEN}\n{JSON}",
        '{"kind":"WindowBlur","durationMs":1500}',
        "200",
        '"strikes" is 0. Short blurs are kept for the dispute trail but never punished.',
        f"""curl -i -X POST {BASE}/api/attempts/{{attemptId}}/proctor-events -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"kind":"WindowBlur","durationMs":1500}}'"""),

    api("API-EXAM-05",
        "A blur longer than 2 seconds is strike 1 and only warns.",
        "Blur lebih dari 2 detik adalah peringatan pertama dan hanya memberi peringatan.",
        "Run UI-EXAM-02 first, with no strikes yet on that attempt.",
        "Jalankan UI-EXAM-02 dulu, belum ada peringatan pada percobaan itu.",
        f"POST {BASE}/api/attempts/{{attemptId}}/proctor-events", f"{TOKEN}\n{JSON}",
        '{"kind":"WindowBlur","durationMs":3000}',
        "200",
        '"strikes" is 1 and "action" is "warn". The attempt is still open.',
        f"""curl -i -X POST {BASE}/api/attempts/{{attemptId}}/proctor-events -H '{TOKEN}' -H '{JSON}' \\
  -d '{{"kind":"WindowBlur","durationMs":3000}}'"""),

    api("API-EXAM-06",
        "Listening audio is refused once its play allowance is used up.",
        "Audio Listening ditolak setelah jatah pemutarannya habis.",
        "Run UI-EXAM-02 first and note a Listening questionId. The play limit is 1.",
        "Jalankan UI-EXAM-02 dulu dan catat questionId Listening. Batas pemutaran 1.",
        f"GET {BASE}/api/attempts/{{attemptId}}/audio/{{questionId}}", TOKEN, "none",
        "409",
        "FIRST call returns 200 with a signed \"url\". SECOND call returns 409. Record the "
        "SECOND call's result here.",
        f"curl -i {BASE}/api/attempts/{{attemptId}}/audio/{{questionId}} -H '{TOKEN}'   # run twice"),

    api("API-EXAM-07",
        "Another learner cannot read someone else's exam sitting.",
        "Pelajar lain tidak dapat membaca pengerjaan tes orang lain.",
        "Two learner accounts. Use learner B's token against learner A's attemptId.",
        "Dua akun pelajar. Pakai token pelajar B terhadap attemptId pelajar A.",
        f"GET {BASE}/api/attempts/{{attemptId_A}}/state", f"{TOKEN} (learner B)", "none",
        "403",
        "No exam data is returned. Report immediately if this returns 200.",
        f"curl -i {BASE}/api/attempts/{{attemptId_A}}/state -H '{TOKEN}'"),

    # ---- certificate ----
    api("API-CERT-01",
        "An unknown certificate code returns a normal answer saying it is not valid.",
        "Kode sertifikat tidak dikenal mengembalikan jawaban normal yang menyatakan tidak sah.",
        "None. No sign-in needed.",
        "Tidak ada. Tidak perlu masuk.",
        f"GET {BASE}/api/program-certificates/verify/NOT-A-REAL-CODE", "none", "none",
        "200",
        '"valid" is false. No name, no score. It is 200 on purpose, not 404 — the page must '
        "still render a normal 'not found' message.",
        f"curl -s {BASE}/api/program-certificates/verify/NOT-A-REAL-CODE"),

    api("API-CERT-02",
        "A real certificate verifies publicly and always carries the prediction disclaimer.",
        "Sertifikat asli terverifikasi publik dan selalu memuat penafian prediksi.",
        "Run UI-EXAM-05 first to finish an exam and get a certificate code from "
        "GET /api/me/program-certificates.",
        "Jalankan UI-EXAM-05 dulu untuk menyelesaikan tes dan ambil kode sertifikat dari "
        "GET /api/me/program-certificates.",
        f"GET {BASE}/api/program-certificates/verify/{{code}}", "none", "none",
        "200",
        '"valid" is true, and "disclaimer" contains "BUKAN skor TOEFL resmi".',
        f"curl -s {BASE}/api/program-certificates/verify/{{code}}"),
]


# ---------------------------------------------------------------- new UI cases
def ui(idv, feat, en, idn, pre_en, pre_id, st_en, st_id, ex_en, ex_id):
    return dict(zip(UI_COLS, [idv, feat, en, idn, pre_en, pre_id, st_en, st_id,
                              ex_en, ex_id, "", "", ""]))


NEW_UI = [
    ui("E-05", "Session test / Tes sesi",
       "A failed session test can be retaken.",
       "Tes sesi yang gagal dapat diulang.",
       "Sign in as siswa2@test.local — the account RESERVED for this case. It is enrolled and has "
       "never taken the test. Do NOT use siswa@test.local: it has already passed, and passing is "
       "permanent, so the retry control is correctly hidden there.",
       "Masuk sebagai siswa2@test.local — akun KHUSUS untuk kasus ini. Sudah terdaftar dan belum "
       "pernah mengerjakan tes. JANGAN pakai siswa@test.local: akun itu sudah lulus, dan kelulusan "
       "bersifat permanen, sehingga tombol ulang memang tidak ditampilkan di sana.",
       "1. Open session 1 and watch the video past ~90% so the test unlocks.\n"
       "2. Answer all 15 questions DELIBERATELY WRONG — pick the first choice every time.\n"
       "3. Click \"Kirim jawaban\".\n"
       "4. Without leaving the page, look at the bottom right of the test panel.",
       "1. Buka sesi 1 dan tonton video sampai lewat ~90% agar tes terbuka.\n"
       "2. Jawab ke-15 soal SENGAJA SALAH — pilih opsi pertama setiap kali.\n"
       "3. Klik \"Kirim jawaban\".\n"
       "4. Tanpa meninggalkan halaman, lihat kanan bawah panel tes.",
       "A button reads \"Coba lagi\" and is clickable. It does NOT read \"Batas percobaan "
       "tercapai\" — this test has no attempt limit. Clicking it clears the answers and reopens "
       "the 15 questions.",
       "Muncul tombol \"Coba lagi\" yang dapat diklik. TIDAK bertuliskan \"Batas percobaan "
       "tercapai\" — tes ini tanpa batas percobaan. Mengkliknya mengosongkan jawaban dan membuka "
       "kembali ke-15 soal."),

    ui("F-01", "Listening audio / Audio listening",
       "Listening questions in a session test can be played.",
       "Soal Listening pada tes sesi dapat diputar.",
       "The session 1 test is open (watch the video past ~90% first). Questions 1-5 are Listening "
       "and each has its own recording.",
       "Tes sesi 1 terbuka (tonton video lewat ~90% dulu). Soal 1-5 adalah Listening dan "
       "masing-masing punya rekaman sendiri.",
       "1. Open the session 1 test and look at question 1.\n"
       "2. Click \"▶ Putar audio\".\n"
       "3. Wait for the player to appear and listen for about 10 seconds.\n"
       "4. Check questions 2 to 5 also have the button, and question 6 onward does not.",
       "1. Buka tes sesi 1 dan lihat soal 1.\n"
       "2. Klik \"▶ Putar audio\".\n"
       "3. Tunggu pemutar muncul lalu dengarkan sekitar 10 detik.\n"
       "4. Periksa soal 2 sampai 5 juga punya tombolnya, dan soal 6 ke atas tidak.",
       "An audio player appears and plays a voice reading a short conversation, matching the "
       "question on screen. Only Listening questions show the button. (A robotic voice is expected "
       "— see the Info sheet.) You may replay as often as you like: a session test has no play "
       "limit, unlike the final exam.",
       "Muncul pemutar audio yang memutar suara membacakan percakapan singkat, sesuai soal di "
       "layar. Hanya soal Listening yang punya tombolnya. (Suara robot itu wajar — lihat sheet "
       "Info.) Anda boleh memutar ulang sesering mungkin: tes sesi tanpa batas pemutaran, berbeda "
       "dengan tes akhir."),

    ui("UI-EXAM-01", "Final assessment / Tes akhir",
       "The final assessment is locked until every earlier session is complete.",
       "Tes akhir terkunci sampai semua sesi sebelumnya selesai.",
       "Signed in as siswa@test.local with ACTIVE access, and at least one earlier session "
       "not yet complete.",
       "Masuk sebagai siswa@test.local dengan akses AKTIF, dan minimal satu sesi sebelumnya "
       "belum selesai.",
       f"1. Open {BASE}/app/dashboard and open the programme.\n"
       "2. Find the last session in the list, named \"Tes Akhir: Simulasi TOEFL ITP\".\n"
       "3. Try to open it.",
       f"1. Buka {BASE}/app/dashboard lalu buka programnya.\n"
       "2. Cari sesi terakhir bernama \"Tes Akhir: Simulasi TOEFL ITP\".\n"
       "3. Coba buka sesi itu.",
       "The session is shown as locked and does not open. No exam starts.",
       "Sesi ditampilkan terkunci dan tidak terbuka. Tidak ada tes yang dimulai."),

    ui("UI-EXAM-02", "Final assessment / Tes akhir",
       "The exam starts on Listening, with a countdown and an audio button.",
       "Tes dimulai pada Listening, dengan hitung mundur dan tombol audio.",
       "Every earlier session complete. TAKES 5 MINUTES TO SET UP. This case STARTS the "
       "learner's only attempt — see UI-EXAM-06 before you run it.",
       "Semua sesi sebelumnya selesai. BUTUH 5 MENIT PERSIAPAN. Kasus ini MEMULAI satu-satunya "
       "percobaan pelajar — baca UI-EXAM-06 sebelum menjalankannya.",
       "1. Open the final session.\n2. Read the notice on the intro screen.\n"
       "3. Click the start button.\n4. Look at the top of the exam screen.",
       "1. Buka sesi terakhir.\n2. Baca pemberitahuan di layar pembuka.\n"
       "3. Klik tombol mulai.\n4. Lihat bagian atas layar tes.",
       "The intro screen states the score is an INVERTA prediction and not an official ETS "
       "TOEFL result. After starting, the screen shows Listening questions, a countdown that "
       "is ticking down, and a button to play the audio.",
       "Layar pembuka menyatakan skor adalah prediksi INVERTA dan bukan skor TOEFL resmi dari "
       "ETS. Setelah dimulai, layar menampilkan soal Listening, hitung mundur yang berjalan, "
       "dan tombol untuk memutar audio."),

    ui("UI-EXAM-03", "Final assessment / Tes akhir",
       "The audio can only be played the allowed number of times.",
       "Audio hanya dapat diputar sebanyak jatah yang diizinkan.",
       "Run UI-EXAM-02 first and stay on the Listening section.",
       "Jalankan UI-EXAM-02 dulu dan tetap di bagian Listening.",
       "1. Read the play button — it shows how many plays are left.\n"
       "2. Click it once and let the clip finish.\n3. Read the button again and click it.",
       "1. Baca tombol putar — tertulis sisa jatah pemutaran.\n"
       "2. Klik sekali dan biarkan klip selesai.\n3. Baca lagi tombolnya dan klik.",
       "After the allowance is used the button shows no plays left and clicking it does not "
       "start the audio again.",
       "Setelah jatah habis, tombol menunjukkan sisa nol dan mengkliknya tidak memutar audio "
       "lagi."),

    ui("UI-EXAM-04", "Final assessment / Tes akhir",
       "Moving to the next section is one-way and the screen says so before you do it.",
       "Pindah ke bagian berikutnya bersifat satu arah dan layar memberitahukannya lebih dulu.",
       "Run UI-EXAM-02 first. IRREVERSIBLE: you cannot return to Listening afterwards.",
       "Jalankan UI-EXAM-02 dulu. TIDAK DAPAT DIBATALKAN: Anda tidak bisa kembali ke Listening.",
       "1. Read the note under the next-section button.\n2. Click the next-section button.\n"
       "3. Look for any way back to Listening — a button, a link, a tab, the browser Back button.",
       "1. Baca catatan di bawah tombol bagian berikutnya.\n2. Klik tombol bagian berikutnya.\n"
       "3. Cari cara kembali ke Listening — tombol, tautan, tab, atau tombol Back peramban.",
       "The note warns you cannot come back. After clicking, Structure questions are shown and "
       "there is no way at all to return to Listening.",
       "Catatan memperingatkan Anda tidak bisa kembali. Setelah diklik, soal Structure "
       "ditampilkan dan tidak ada cara apa pun kembali ke Listening."),

    ui("UI-EXAM-05", "Final assessment / Tes akhir",
       "Finishing the exam shows a predicted score and issues a certificate.",
       "Menyelesaikan tes menampilkan skor prediksi dan menerbitkan sertifikat.",
       "Run UI-EXAM-04 twice so you are on Reading. IRREVERSIBLE: this uses up the account's "
       "only attempt. Run UI-EXAM-06 after this, never before.",
       "Jalankan UI-EXAM-04 dua kali sampai Anda di Reading. TIDAK DAPAT DIBATALKAN: ini "
       "menghabiskan satu-satunya percobaan akun. Jalankan UI-EXAM-06 setelah ini, jangan "
       "sebelumnya.",
       "1. Answer some Reading questions.\n2. Click the finish button.\n3. Read the result "
       "screen from top to bottom.",
       "1. Jawab beberapa soal Reading.\n2. Klik tombol selesai.\n3. Baca layar hasil dari "
       "atas sampai bawah.",
       "A predicted TOEFL score between 310 and 677 is shown. The screen states plainly that "
       "this is an INVERTA prediction and NOT an official ETS TOEFL result. A link to the "
       "certificate is offered.",
       "Skor prediksi TOEFL antara 310 dan 677 ditampilkan. Layar menyatakan dengan jelas ini "
       "prediksi INVERTA dan BUKAN skor TOEFL resmi dari ETS. Tersedia tautan ke sertifikat."),

    ui("UI-EXAM-06", "Final assessment / Tes akhir",
       "A second attempt is refused once the single attempt is used.",
       "Percobaan kedua ditolak setelah satu-satunya percobaan terpakai.",
       "Run UI-EXAM-05 first. Run this AFTER it, never before.",
       "Jalankan UI-EXAM-05 dulu. Jalankan ini SESUDAHNYA, jangan sebelumnya.",
       "1. Open the final session again.\n2. Look at the button on the intro screen.",
       "1. Buka lagi sesi terakhir.\n2. Lihat tombol di layar pembuka.",
       "The button says the attempt limit is reached and cannot be clicked. The exam does not "
       "restart.",
       "Tombol menyatakan batas percobaan tercapai dan tidak bisa diklik. Tes tidak dimulai "
       "ulang."),

    ui("UI-EXAM-07", "Final assessment / Tes akhir",
       "The certificate can be verified by anyone, without signing in.",
       "Sertifikat dapat diverifikasi siapa pun, tanpa masuk.",
       "Run UI-EXAM-05 first and open the certificate to copy its verification code.",
       "Jalankan UI-EXAM-05 dulu dan buka sertifikat untuk menyalin kode verifikasinya.",
       "1. Open a private/incognito window so you are signed out.\n"
       f"2. Go to {BASE}/verify/<the code you copied>.\n3. Read the page.",
       "1. Buka jendela penyamaran agar Anda dalam keadaan keluar.\n"
       f"2. Buka {BASE}/verify/<kode yang Anda salin>.\n3. Baca halamannya.",
       "The page loads without signing in, shows the learner name and the predicted score, and "
       "states it is an INVERTA prediction, not an official ETS TOEFL result.",
       "Halaman terbuka tanpa masuk, menampilkan nama pelajar dan skor prediksi, serta "
       "menyatakan ini prediksi INVERTA, bukan skor TOEFL resmi dari ETS."),

    ui("UI-EXAM-08", "Final assessment / Tes akhir",
       "An unknown certificate code shows a not-found message, not an error page.",
       "Kode sertifikat tidak dikenal menampilkan pesan tidak ditemukan, bukan halaman error.",
       "None. You can be signed out.",
       "Tidak ada. Anda boleh dalam keadaan keluar.",
       f"1. Go to {BASE}/verify/NOTAREALCODE123.\n2. Read the page.",
       f"1. Buka {BASE}/verify/NOTAREALCODE123.\n2. Baca halamannya.",
       "The page loads normally and says the certificate was not found. It does NOT show a "
       "certificate, and it is not a browser error page.",
       "Halaman terbuka normal dan menyatakan sertifikat tidak ditemukan. TIDAK menampilkan "
       "sertifikat, dan bukan halaman error peramban."),

    ui("UI-ADM-08", "Admin — Learners / Peserta",
       "A revoked learner can be restored from the enrollment row.",
       "Pelajar yang dicabut dapat dipulihkan dari baris pendaftarannya.",
       "Signed in as admin, and siswa@test.local has been revoked (run I-02 first).",
       "Masuk sebagai admin, dan siswa@test.local sudah dicabut (jalankan I-02 dulu).",
       f"1. Go to {BASE}/admin/enrollments and find siswa@test.local. Its status is \"Dicabut\".\n"
       "2. Click \"Pulihkan\" on that row and confirm.\n3. Read the status column.",
       f"1. Buka {BASE}/admin/enrollments lalu cari siswa@test.local. Statusnya \"Dicabut\".\n"
       "2. Klik \"Pulihkan\" pada baris itu lalu konfirmasi.\n3. Baca kolom status.",
       "The status returns to \"Aktif\". Only ONE row exists for that learner — restoring must "
       "not create a second enrollment. Their earlier progress is still there.",
       "Status kembali menjadi \"Aktif\". Hanya ADA SATU baris untuk pelajar itu — memulihkan "
       "tidak boleh membuat pendaftaran kedua. Progres sebelumnya masih ada."),

    ui("UI-IMP-01", "Admin — Bulk import / Impor massal",
       "The question import template can be downloaded and it opens in Excel.",
       "Template impor soal dapat diunduh dan terbuka di Excel.",
       "Signed in as admin.",
       "Masuk sebagai admin.",
       f"1. Go to {BASE}/admin/questions/import.\n2. Click the template download button.\n"
       "3. Open the downloaded file.",
       f"1. Buka {BASE}/admin/questions/import.\n2. Klik tombol unduh template.\n"
       "3. Buka berkas yang terunduh.",
       "An .xlsx file downloads and opens without a repair prompt. It has sheets named "
       "\"Questions\", \"Passages\" and \"Instructions\", and the Questions sheet already has "
       "example rows.",
       "Berkas .xlsx terunduh dan terbuka tanpa permintaan perbaikan. Ada sheet bernama "
       "\"Questions\", \"Passages\" dan \"Instructions\", dan sheet Questions sudah berisi "
       "contoh baris."),

    ui("UI-IMP-02", "Admin — Bulk import / Impor massal",
       "A sheet with a mistake is refused and every mistake is listed with its row and column.",
       "Berkas dengan kesalahan ditolak dan setiap kesalahan didaftar beserta baris dan kolomnya.",
       "Run UI-IMP-01 first. In the template, change the \"section\" cell of the first example "
       "row to \"Speaking\" and save.",
       "Jalankan UI-IMP-01 dulu. Pada template, ubah sel \"section\" baris contoh pertama "
       "menjadi \"Speaking\" lalu simpan.",
       f"1. Go to {BASE}/admin/questions/import.\n2. Choose your edited file.\n"
       "3. Click the check button.\n4. Read the error table.",
       f"1. Buka {BASE}/admin/questions/import.\n2. Pilih berkas yang Anda ubah.\n"
       "3. Klik tombol periksa.\n4. Baca tabel kesalahan.",
       "An error table appears naming the ROW NUMBER and the COLUMN \"section\", and says which "
       "section names are allowed. The import button stays disabled — nothing is saved.",
       "Muncul tabel kesalahan yang menyebutkan NOMOR BARIS dan KOLOM \"section\", serta "
       "menyatakan nama bagian yang diizinkan. Tombol impor tetap nonaktif — tidak ada yang "
       "tersimpan."),
]


# ---------------------------------------------------------------- build
def carried_over_ui():
    """v1.0 cases that were already UI-shaped, kept verbatim so reviewed wording is not lost."""
    src = load_workbook(SRC)["Test Cases"]
    hdr = [c.value for c in src[1]]
    out = []
    for row in src.iter_rows(min_row=2, values_only=True):
        if not row[0] or row[0] in MOVED_TO_API or row[0] in REWRITTEN:
            continue
        if not str(row[0])[1:2] == "-":          # skip the section divider row
            continue
        r = dict(zip(hdr, row))
        out.append(ui(r["ID"], r["Feature / Fitur"], r["Description (EN)"], r["Deskripsi (ID)"],
                      r["Preconditions (EN)"], r["Prasyarat (ID)"], r["Steps (EN)"],
                      r["Langkah (ID)"], r["Expected Result (EN)"],
                      r["Hasil yang Diharapkan (ID)"]))
    return out


def sheet(wb, title, cols, rows, widths):
    ws = wb.create_sheet(title)
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c)
        cell.value = name
        cell.fill = HEAD
        cell.font = HEADFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=c)
            cell.value = row.get(name, "")
            cell.alignment = WRAP
        ws.row_dimensions[r].height = 92
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    return ws


def info_sheet(wb, n_api, n_ui):
    ws = wb.create_sheet("Info", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 96
    rows = [
        ("INVERTA — User Acceptance Test (UAT)", ""),
        ("Uji Penerimaan Pengguna", ""),
        ("", ""),
        ("Version / Versi", "2.0 — split into API and UI sheets"),
        ("Environment / Lingkungan", BASE),
        ("Date / Tanggal", ""),
        ("Tester / Penguji", ""),
        ("", ""),
        ("1. WHICH SHEET DO I USE?", ""),
        ("UI Test Cases", "You click through the website in a browser. No tools needed."),
        ("API Test Cases", "You send requests with a REST client (Postman/Insomnia) or paste "
                           "the ready-made curl command from the last column into a terminal."),
        ("", "Do the UI sheet first. The API sheet checks rules you cannot see on screen — "
             "mostly security rules and exact error codes."),
        ("", ""),
        ("2. ACCOUNTS / AKUN", ""),
        ("Learner / Siswa", "siswa@test.local"),
        ("Learner for retakes", "siswa2@test.local — RESERVED for E-05. Do NOT pass its session "
                                "test. Passing is permanent and attempts are never deleted, so "
                                "once passed the retry control is correctly hidden for ever and "
                                "E-05 can never be run on that account again."),
        ("Admin", "admin@academy.local"),
        ("", "Passwords are deliberately not written here. Ask your lead once, then keep them "
             "yourself. Kata sandi sengaja tidak ditulis di sini. Tanyakan ke lead Anda."),
        ("", ""),
        ("3. HOW TO GET A TOKEN (API sheet only)", ""),
        ("Step 1", f"curl -s -X POST {BASE}/api/auth/login -H 'content-type: application/json' "
                   "-d '{\"email\":\"siswa@test.local\",\"password\":\"<ask your lead>\"}'"),
        ("Step 2", "Copy the value of \"accessToken\" from the reply. Paste it wherever a case "
                   "shows <token>."),
        ("Step 3", "The token stops working after 15 minutes. Repeat step 1 to get a new one."),
        ("Admin token", "Same two steps, using admin@academy.local."),
        ("", ""),
        ("4. HOW TO RUN A CASE", ""),
        ("", "• Read Preconditions first. If they are not true, make them true before you start."),
        ("", "  Baca Prasyarat dulu. Kalau belum terpenuhi, penuhi dulu sebelum mulai."),
        ("", "• Follow the steps exactly, in order. Do not skip or combine steps."),
        ("", "  Ikuti langkah persis, berurutan. Jangan melewati atau menggabungkan langkah."),
        ("", "• Compare what you see with the Expected Result. They must match exactly."),
        ("", "  Bandingkan yang Anda lihat dengan Hasil yang Diharapkan. Harus sama persis."),
        ("", "• Fill Result with Pass or Fail. If Fail, write EXACTLY what you saw in Notes."),
        ("", "  Isi Hasil dengan Pass atau Fail. Kalau Fail, tulis persis apa yang Anda lihat."),
        ("", "• A case you cannot start is Blocked, not Fail. Say why in Notes."),
        ("", "  Kasus yang tidak bisa dimulai adalah Blocked, bukan Fail. Tulis alasannya."),
        ("", "• Text in \"quotes\" is exactly what appears on screen, in Indonesian."),
        ("", "  Teks dalam \"tanda kutip\" persis seperti di layar, dalam Bahasa Indonesia."),
        ("", "• Run cases in ID order unless a case says otherwise. Some cases are one-shot and "
             "say so — read those before starting them."),
        ("", "  Jalankan sesuai urutan ID kecuali kasus menyebutkan lain. Beberapa kasus hanya "
             "bisa sekali dan menyebutkannya — baca dulu sebelum menjalankannya."),
        ("", ""),
        ("5. NOT PART OF THIS BUILD — DO NOT RAISE BUGS", ""),
        ("Bukan bagian dari build ini — jangan laporkan sebagai bug", ""),
        ("✗ Real emails", "Emails are written to the server log, not sent. You will never "
                          "receive one. Email verification therefore cannot be completed by a "
                          "real inbox."),
        ("✗ Google sign-in", "The \"Lanjutkan dengan Google\" button is visible but was never "
                             "connected. It will not work."),
        ("✗ Real payment", "Payment uses a built-in simulator with a succeed and a fail button. "
                           "No real money, no real gateway."),
        ("✗ Video content", "Every session plays the same sample clip on purpose."),
        ("✗ Exam question wording", "The final exam holds 140 PLACEHOLDER questions. Every one "
                                    "starts with \"[CONTOH]\". They exist so the exam can be "
                                    "tested; they are not real TOEFL content. Do not report "
                                    "their wording, difficulty or answers as bugs."),
        ("✗ Listening audio quality", "The sample audio is computer-generated on purpose. A "
                                      "robotic voice is expected."),
        ("✗ Section time-out message", "When a section's time runs out the exam moves on "
                                       "silently. There is no \"waktu habis\" message in this "
                                       "build. Do not report the absence as a bug."),
        ("", ""),
        ("6. WHAT IS IN THIS PACK", ""),
        ("API test cases", n_api),
        ("UI test cases", n_ui),
        ("TOTAL", n_api + n_ui),
        ("", ""),
        ("", "Cases marked SECURITY must be reported to the lead IMMEDIATELY if they fail — "
             "do not wait until the end of the run."),
        ("", "Kasus bertanda SECURITY harus dilaporkan ke lead SEGERA jika gagal — jangan "
             "menunggu sampai pengujian selesai."),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ws.cell(row=r, column=1).value = a
        ws.cell(row=r, column=2).value = b
        ws.cell(row=r, column=1).alignment = WRAP
        ws.cell(row=r, column=2).alignment = WRAP
        if a and not b and a[0].isdigit() or a.startswith(("INVERTA", "Uji", "Bukan")):
            ws.cell(row=r, column=1).font = Font(bold=True)
    return ws


def main():
    ui_rows = NEW_UI + carried_over_ui()
    wb = Workbook()
    wb.remove(wb.active)

    info_sheet(wb, len(API_CASES), len(ui_rows))
    sheet(wb, "API Test Cases", API_COLS, API_CASES,
          [14, 34, 34, 34, 34, 40, 26, 34, 12, 44, 52, 9, 11, 30])
    sheet(wb, "UI Test Cases", UI_COLS, ui_rows,
          [12, 26, 34, 34, 32, 32, 40, 40, 40, 40, 9, 11, 30])

    # Answer key carried over untouched — QA needs it to pass a test on purpose.
    src = load_workbook(SRC)["Answer key"]
    dst = wb.create_sheet("Answer key")
    for r, row in enumerate(src.iter_rows(values_only=True), start=1):
        for c, v in enumerate(row, start=1):
            dst.cell(row=r, column=c).value = v
    dst.column_dimensions["A"].width = 12
    dst.column_dimensions["B"].width = 70
    dst.column_dimensions["C"].width = 40

    wb.save(OUT)
    print(f"wrote {OUT}: {len(API_CASES)} API + {len(ui_rows)} UI = "
          f"{len(API_CASES) + len(ui_rows)} cases")


if __name__ == "__main__":
    main()
